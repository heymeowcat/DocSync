import os
import json
from git import Repo
from transformers import T5ForConditionalGeneration, T5Tokenizer
from docx import Document
import difflib
import logging
import torch

# Set up logging
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

# Load T5 model for text generation
model_name = "google/flan-t5-small"
tokenizer = T5Tokenizer.from_pretrained(model_name)
model = T5ForConditionalGeneration.from_pretrained(model_name)

# Use MPS (Metal Performance Shaders) backend if available
device = "mps" if torch.backends.mps.is_available() else "cpu"
model.to(device)

def generate_lld_content(code_changes, commit_info):
    """Generate LLD content using T5 model"""
    prompt = f"""Generate a Low-Level Design (LLD) section based on these code changes:
    Commit Message: {commit_info['message']}
    Author: {commit_info['author']}
    File Changes: {json.dumps(code_changes)}
    
    Provide detailed design insights, focusing on:
    - Module purpose
    - Responsibilities
    - Key algorithmic changes
    - Potential performance implications
    """
    
    # Tokenize the prompt
    inputs = tokenizer(prompt, return_tensors="pt", max_length=512, truncation=True).to(device)
    
    # Generate LLD content
    outputs = model.generate(**inputs, max_length=300, num_return_sequences=1)
    lld_text = tokenizer.decode(outputs[0], skip_special_tokens=True)
    
    return lld_text

def get_commit_hashes(repo_path):
    repo = Repo(repo_path)
    return [str(commit) for commit in list(repo.iter_commits())]

def generate_detailed_commit_info(repo_path, old_commit, new_commit, lld_path=None):
    repo = Repo(repo_path)
    
    # Get commits
    old_commit_obj = repo.commit(old_commit)
    new_commit_obj = repo.commit(new_commit)
    
    # Detailed commit entries
    detailed_entries = []
    
    # Commits between old and new
    commits = list(repo.iter_commits(f"{old_commit}..{new_commit}"))
    
    for commit in commits:
        # Basic commit information
        entry = {
            'commit_hash': str(commit),
            'message': commit.message,
            'author': commit.author.name,
            'email': commit.author.email,
            'date': commit.committed_datetime.strftime('%Y-%m-%d %H:%M:%S')
        }
        
        # Collect file changes
        file_changes = []
        diffs = []
        
        for d in commit.diff(commit.parents[0] if commit.parents else None):
            change_type = {
                'A': 'Added',
                'D': 'Deleted',
                'M': 'Modified',
                'R': 'Renamed'
            }.get(d.change_type, 'Unknown')
            
            file_change = {
                'file': d.a_path,
                'change_type': change_type
            }
            file_changes.append(file_change)
            
            # Get diff content for modified files
            if change_type == 'Modified':
                try:
                    blob_old = old_commit_obj.tree[d.a_path]
                    blob_new = new_commit_obj.tree[d.a_path]
                    
                    old_content = blob_old.data_stream.read().decode('utf-8', errors='ignore').splitlines()
                    new_content = blob_new.data_stream.read().decode('utf-8', errors='ignore').splitlines()
                    
                    diff_lines = list(difflib.unified_diff(
                        old_content, 
                        new_content, 
                        fromfile=f'a/{d.a_path}', 
                        tofile=f'b/{d.a_path}'
                    ))
                    
                    # Generate LLD for this specific change
                    lld_section = generate_lld_content(file_change, entry)
                    
                    diffs.append({
                        'file': d.a_path,
                        'diff_content': diff_lines,
                        'lld_section': lld_section
                    })
                except KeyError as e:
                    logger.warning(f"File {d.a_path} not found in commit {commit.hexsha}: {e}")
                except Exception as e:
                    logger.error(f"Error processing {d.a_path}: {e}")
        
        entry['file_changes'] = file_changes
        entry['diffs'] = diffs
        detailed_entries.append(entry)
    
    # Create LLD document
    if not lld_path:
        lld_path = os.path.join(repo_path, 'project_lld.docx')
    
    doc = Document()
    doc.add_heading('Low-Level Design (LLD) Document', 0)
    
    # Introduction
    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph('1.1 Purpose\nExplain the purpose of the LLD document.')
    doc.add_paragraph('1.2 Scope\nDefine the scope of the LLD, including what aspects of the system will be covered.')
    doc.add_paragraph('1.3 Audience\nIdentify the intended audience for this document.')
    doc.add_paragraph('1.4 References\nList any documents, standards, or guidelines that are referenced in this document.')
    
    # System Overview
    doc.add_heading('2. System Overview', level=1)
    doc.add_paragraph('2.1 System Description\nProvide a brief overview of the system, including its purpose and functionality.')
    doc.add_paragraph('2.2 System Context\nDescribe how the system fits into the larger business or technical environment.')
    
    # Detailed Design
    doc.add_heading('3. Detailed Design', level=1)
    doc.add_paragraph('3.1 Module Descriptions\nDescribe each module in detail, including its purpose and functionality.')
    doc.add_paragraph('3.2 Class Diagrams\nProvide class diagrams that detail the structure of each module.')
    doc.add_paragraph('3.3 Sequence Diagrams\nInclude sequence diagrams that illustrate the interactions between objects for various scenarios.')
    doc.add_paragraph('3.4 State Diagrams\nProvide state diagrams that describe the state changes of important objects.')
    doc.add_paragraph('3.5 Activity Diagrams\nInclude activity diagrams to show the workflow and operations within the system.')
    
    # Data Design
    doc.add_heading('4. Data Design', level=1)
    doc.add_paragraph('4.1 Data Structures\nDescribe the data structures used in the system, including data types and relationships.')
    doc.add_paragraph('4.2 Database Design\nDetail the database design, including tables, columns, keys, and relationships.')
    doc.add_paragraph('4.3 Data Flow\nDescribe how data moves through the system at a detailed level.')
    
    # Interface Design
    doc.add_heading('5. Interface Design', level=1)
    doc.add_paragraph('5.1 User Interface\nProvide detailed designs for the user interface, including screen layouts, elements, and behaviors.')
    doc.add_paragraph('5.2 External Interfaces\nDescribe the design of external interfaces, such as APIs, including request and response formats.')
    doc.add_paragraph('5.3 Interface Contracts\nDetail the contracts for each interface, including input and output specifications.')
    
    # Algorithm Design
    doc.add_heading('6. Algorithm Design', level=1)
    doc.add_paragraph('6.1 Algorithms\nDescribe the algorithms used in the system, including pseudocode or flowcharts.')
    doc.add_paragraph('6.2 Complexity Analysis\nProvide an analysis of the complexity of each algorithm.')
    
    # Security Design
    doc.add_heading('7. Security Design', level=1)
    doc.add_paragraph('7.1 Security Measures\nDetail the security measures implemented at a detailed level, such as encryption methods and access controls.')
    doc.add_paragraph('7.2 Authentication and Authorization\nDescribe the mechanisms for authentication and authorization in detail.')
    doc.add_paragraph('7.3 Data Protection\nDetail how data is protected at rest and in transit.')
    
    # Error Handling and Logging
    doc.add_heading('8. Error Handling and Logging', level=1)
    doc.add_paragraph('8.1 Error Handling\nDescribe how errors are handled at a detailed level, including specific error messages and recovery procedures.')
    doc.add_paragraph('8.2 Logging\nDetail the logging strategy, including log formats, levels, and storage.')
    
    # Performance Considerations
    doc.add_heading('9. Performance Considerations', level=1)
    doc.add_paragraph('9.1 Performance Optimization\nDetail the strategies for optimizing performance, including specific techniques used.')
    doc.add_paragraph('9.2 Load Handling\nDescribe how the system handles load, including load balancing and resource management.')
    
    # Testing and Validation
    doc.add_heading('10. Testing and Validation', level=1)
    doc.add_paragraph('10.1 Unit Testing\nDescribe the unit testing strategy, including test cases and tools used.')
    doc.add_paragraph('10.2 Integration Testing\nDetail the integration testing approach, including scenarios and tools.')
    doc.add_paragraph('10.3 Validation\nDescribe how the system will be validated against requirements.')
    
    # Deployment Considerations
    doc.add_heading('11. Deployment Considerations', level=1)
    doc.add_paragraph('11.1 Deployment Architecture\nDetail the deployment architecture, including environments and configurations.')
    doc.add_paragraph('11.2 Deployment Process\nDescribe the process for deploying the system, including steps and tools.')
    
    # Assumptions and Dependencies
    doc.add_heading('12. Assumptions and Dependencies', level=1)
    doc.add_paragraph('12.1 Assumptions\nList any assumptions made during the design process.')
    doc.add_paragraph('12.2 Dependencies\nIdentify any dependencies on other systems, components, or technologies.')
    
    # Appendix
    doc.add_heading('13. Appendix', level=1)
    doc.add_paragraph('13.1 Glossary\nProvide a glossary of terms used in the document.')
    doc.add_paragraph('13.2 Acronyms\nList and define any acronyms used in the document.')
    doc.add_paragraph('13.3 Document History\nInclude a version history of the document, noting changes and updates.')
    
    for entry in detailed_entries:
        doc.add_heading(f'Commit: {entry["commit_hash"]}', level=1)
        doc.add_paragraph(f'Message: {entry["message"]}')
        doc.add_paragraph(f'Author: {entry["author"]} <{entry["email"]}>')
        doc.add_paragraph(f'Date: {entry["date"]}')
        
        for diff in entry.get('diffs', []):
            doc.add_heading(f'File: {diff["file"]}', level=2)
            doc.add_paragraph('Changes:')
            doc.add_paragraph('\n'.join(diff['diff_content']))
            
            doc.add_heading('LLD Insights', level=3)
            doc.add_paragraph(diff['lld_section'])
    
    doc.save(lld_path)
    
    return {
        'detailed_entries': detailed_entries,
        'lld_path': lld_path
    }

def main():
    repo_path = input("Enter the path to the Git repository: ")
    repo = Repo(repo_path)
    
    branches = [branch.name for branch in repo.branches]
    print("Available branches:")
    for i, branch in enumerate(branches):
        print(f"{i + 1}. {branch}")
    
    branch_index = int(input("Select a branch by number: ")) - 1
    selected_branch = branches[branch_index]
    
    commits = list(repo.iter_commits(selected_branch))
    print("Available commits:")
    for i, commit in enumerate(commits):
        print(f"{i + 1}. {commit.hexsha} - {commit.message.strip()}")
    
    from_commit_index = int(input("Select the 'from' commit by number: ")) - 1
    to_commit_index = int(input("Select the 'to' commit by number: ")) - 1
    
    from_commit = commits[from_commit_index].hexsha
    to_commit = commits[to_commit_index].hexsha
    
    generate_detailed_commit_info(repo_path, from_commit, to_commit)
    print("LLD document generated successfully.")

if __name__ == '__main__':
    main()